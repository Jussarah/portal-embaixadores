"""Lógica do Portal de Embaixadores BS LATAM: correção de links, análise e relatórios."""
import io
import json
import re
import urllib.parse
import urllib.request
from datetime import datetime
from urllib.parse import parse_qs, urlparse

import pandas as pd

PLATAFORMAS = ["Instagram", "TikTok", "YouTube", "Facebook", "Kwai"]
COLS_VIDEOS = ["Link", "Embaixador", "Rede", "Data", "Visualizações", "Curtidas", "Salvamentos"]
COLS_NUM = ["Visualizações", "Curtidas", "Salvamentos"]

USER_RE = re.compile(r"^[A-Za-z0-9._]{2,30}$")

ST_OK = "✅ OK"
ST_FIX = "🛠️ Corrigido"
ST_CHECK = "⚠️ Verificar"
ST_BAD = "❌ Inválido"
ST_DUP = "🔁 Duplicado"


# --------------------------------------------------------------------------- #
# Utilidades de URL
# --------------------------------------------------------------------------- #
def _is(host, *domains):
    return any(host == d or host.endswith("." + d) for d in domains)


def _plataforma_do_host(host):
    if _is(host, "instagram.com", "instagr.am"):
        return "Instagram"
    if _is(host, "tiktok.com"):
        return "TikTok"
    if _is(host, "youtube.com", "youtu.be"):
        return "YouTube"
    if _is(host, "facebook.com", "fb.com", "fb.watch", "fb.me"):
        return "Facebook"
    if _is(host, "kwai.com", "kw.ai", "kwai-video.com"):
        return "Kwai"
    return None


def _parse(raw):
    s = str(raw or "").strip().strip("<>\"'“”‘’,;()[]")
    if not s:
        return None
    if not re.match(r"^[a-zA-Z][a-zA-Z0-9+.-]*://", s):
        s = "https://" + s.lstrip("/")
    return urlparse(s)


def _host(u):
    host = (u.hostname or "").lower()
    for p in ("www.", "m.", "mobile.", "web.", "mbasic.", "l.", "touch."):
        if host.startswith(p):
            host = host[len(p):]
            break
    return host


def detectar_rede(url):
    u = _parse(url)
    if not u:
        return "Outra"
    return _plataforma_do_host(_host(u)) or "Outra"


# --------------------------------------------------------------------------- #
# Corretor de links de perfil
# --------------------------------------------------------------------------- #
def _res(original, rede, corrigido, status, obs):
    return {
        "Original": original,
        "Rede": rede or "—",
        "Link corrigido": corrigido,
        "Status": status,
        "Observação": obs,
    }


def _final(original, rede, corrigido, notas, u=None):
    igual = original.strip().rstrip("/") == corrigido.rstrip("/")
    if igual:
        return _res(original, rede, corrigido, ST_OK, "Já está no padrão")
    if not notas:
        notas = ["parâmetros de rastreio removidos" if (u is not None and u.query) else "link padronizado"]
    return _res(original, rede, corrigido, ST_FIX, "; ".join(notas))


def _instagram(orig, u, segs):
    reservadas = {"explore", "accounts", "direct", "about", "web", "legal", "privacy", "developer"}
    if not segs:
        return _res(orig, "Instagram", "", ST_BAD, "Link sem usuário")
    first = segs[0].lower()
    if first in ("p", "reel", "reels", "tv"):
        return _res(orig, "Instagram", f"https://www.instagram.com{u.path}", ST_CHECK,
                    "Link de publicação, não identifica o perfil")
    if first == "stories" and len(segs) > 1:
        user = segs[1]
        notas = ["perfil extraído do link de story"]
    elif len(segs) > 1 and segs[1].lower() in ("p", "reel", "tv"):
        user = segs[0]
        notas = ["perfil extraído do link da publicação"]
    elif first in reservadas:
        return _res(orig, "Instagram", "", ST_BAD, "Não é link de perfil")
    else:
        user = segs[0]
        notas = []
    if not USER_RE.match(user):
        return _res(orig, "Instagram", "", ST_BAD, "Nome de usuário inválido")
    return _final(orig, "Instagram", f"https://www.instagram.com/{user}/", notas, u)


def _tiktok(orig, host, u, segs):
    reservadas = {"foryou", "discover", "tag", "music", "search", "explore", "live", "following", "friends"}
    if host.startswith(("vm.", "vt.")) or (segs and segs[0] == "t"):
        return _res(orig, "TikTok", f"https://{host}{u.path}", ST_CHECK,
                    "Link curto: abra no navegador e copie o link do perfil")
    if not segs:
        return _res(orig, "TikTok", "", ST_BAD, "Link sem usuário")
    first = segs[0]
    notas = []
    if first.startswith("@"):
        user = first[1:]
        if len(segs) > 1 and segs[1].lower() in ("video", "photo", "live"):
            notas.append("perfil extraído do link do vídeo")
    elif first.lower() in reservadas:
        return _res(orig, "TikTok", "", ST_BAD, "Não é link de perfil")
    else:
        user = first
        notas.append("@ adicionado ao usuário")
    if not re.match(r"^[A-Za-z0-9._]{2,24}$", user):
        return _res(orig, "TikTok", "", ST_BAD, "Nome de usuário inválido")
    return _final(orig, "TikTok", f"https://www.tiktok.com/@{user}", notas, u)


def _youtube(orig, host, u, segs, q):
    if host == "youtu.be":
        return _res(orig, "YouTube", f"https://youtu.be{u.path}", ST_CHECK, "Link de vídeo, não de canal")
    if not segs:
        return _res(orig, "YouTube", "", ST_BAD, "Link sem canal")
    first = segs[0]
    low = first.lower()
    if low in ("watch", "shorts", "live", "embed", "playlist", "results", "feed") or "v" in q:
        if "v" in q:
            corr = f"https://www.youtube.com/watch?v={q['v'][0]}"
        else:
            corr = f"https://www.youtube.com{u.path}"
        return _res(orig, "YouTube", corr, ST_CHECK, "Link de vídeo ou lista, não de canal")
    if first.startswith("@"):
        notas = ["sufixo da página removido"] if len(segs) > 1 else []
        return _final(orig, "YouTube", f"https://www.youtube.com/{first}", notas, u)
    if low == "channel" and len(segs) > 1:
        return _final(orig, "YouTube", f"https://www.youtube.com/channel/{segs[1]}", [], u)
    if low in ("c", "user") and len(segs) > 1:
        return _final(orig, "YouTube", f"https://www.youtube.com/{low}/{segs[1]}", [], u)
    return _res(orig, "YouTube", f"https://www.youtube.com/@{first}", ST_CHECK,
                "Sem @ no link: confirme se este é o canal")


def _facebook(orig, host, u, segs, q):
    posts = {"watch", "reel", "reels", "share", "photo", "photos", "permalink.php", "story.php", "video",
             "videos", "groups", "events", "hashtag", "stories", "sharer", "sharer.php", "plugins", "login"}
    if host == "fb.watch":
        return _res(orig, "Facebook", f"https://fb.watch{u.path}", ST_CHECK, "Link de vídeo, não de perfil")
    if not segs:
        return _res(orig, "Facebook", "", ST_BAD, "Link sem perfil")
    first = segs[0]
    low = first.lower()
    if low == "profile.php":
        pid = (q.get("id") or [""])[0]
        if pid.isdigit():
            return _final(orig, "Facebook", f"https://www.facebook.com/profile.php?id={pid}", [], u)
        return _res(orig, "Facebook", "", ST_BAD, "profile.php sem ID numérico")
    if low == "people" and len(segs) >= 3:
        return _final(orig, "Facebook", f"https://www.facebook.com/people/{segs[1]}/{segs[2]}/", [], u)
    if low == "pg" and len(segs) > 1:
        first = segs[1]
        low = first.lower()
    if low in posts:
        return _res(orig, "Facebook", f"https://www.facebook.com{u.path}", ST_CHECK,
                    "Link de publicação, vídeo ou grupo, não de perfil")
    if not re.match(r"^[A-Za-z0-9.\-]{5,50}$", first):
        return _res(orig, "Facebook", "", ST_CHECK, "Usuário fora do padrão: confirme o link")
    return _final(orig, "Facebook", f"https://www.facebook.com/{first}", [], u)


def _kwai(orig, host, u, segs):
    if host in ("kw.ai", "k.kwai.com") or (segs and segs[0].lower() in ("s", "sh")):
        return _res(orig, "Kwai", f"https://{host}{u.path}", ST_CHECK,
                    "Link curto: abra no navegador e copie o link do perfil")
    if not segs:
        return _res(orig, "Kwai", "", ST_BAD, "Link sem perfil")
    first = segs[0]
    if first.startswith("@"):
        user = first[1:]
        if not USER_RE.match(user):
            return _res(orig, "Kwai", "", ST_BAD, "Nome de usuário inválido")
        return _final(orig, "Kwai", f"https://www.kwai.com/@{user}", [], u)
    if first.lower() == "profile" and len(segs) > 1:
        return _final(orig, "Kwai", f"https://www.kwai.com/profile/{segs[1]}", [], u)
    if first.lower() in ("video", "photo", "live", "feed"):
        return _res(orig, "Kwai", f"https://www.kwai.com{u.path}", ST_CHECK, "Link de vídeo, não de perfil")
    return _res(orig, "Kwai", f"https://www.kwai.com/@{first}", ST_CHECK,
                "Sem @ no link: confirme abrindo o perfil")


def corrigir_link(raw):
    original = str(raw).strip()
    u = _parse(original)
    if not u or not u.hostname:
        return _res(original, None, "", ST_BAD, "Link vazio ou ilegível")
    host = _host(u)
    rede = _plataforma_do_host(host)
    segs = [s for s in u.path.split("/") if s]
    q = parse_qs(u.query)
    if rede == "Instagram":
        return _instagram(original, u, segs)
    if rede == "TikTok":
        return _tiktok(original, host, u, segs)
    if rede == "YouTube":
        return _youtube(original, host, u, segs, q)
    if rede == "Facebook":
        return _facebook(original, host, u, segs, q)
    if rede == "Kwai":
        return _kwai(original, host, u, segs)
    return _res(original, None, "", ST_BAD, "Rede não reconhecida")


def corrigir_lista(texto):
    """Recebe links separados por espaço, vírgula ou quebra de linha."""
    tokens = [t for t in re.split(r"[\s,;]+", texto or "") if t.strip()]
    vistos, resultados = set(), []
    for t in tokens:
        r = corrigir_link(t)
        chave = r["Link corrigido"].lower().rstrip("/")
        if chave and chave in vistos:
            r["Status"] = ST_DUP
            r["Observação"] = "Mesmo perfil já está na lista"
        elif chave:
            vistos.add(chave)
        resultados.append(r)
    return resultados


# --------------------------------------------------------------------------- #
# YouTube (opcional, com chave gratuita da API do Google)
# --------------------------------------------------------------------------- #
def youtube_video_id(url):
    u = _parse(url)
    if not u:
        return None
    host = _host(u)
    segs = [s for s in u.path.split("/") if s]
    if host == "youtu.be" and segs:
        return segs[0]
    q = parse_qs(u.query)
    if "v" in q:
        return q["v"][0]
    if len(segs) > 1 and segs[0] in ("shorts", "live", "embed"):
        return segs[1]
    return None


def buscar_youtube(video_id, api_key):
    params = urllib.parse.urlencode({"part": "statistics,snippet", "id": video_id, "key": api_key})
    with urllib.request.urlopen(f"https://www.googleapis.com/youtube/v3/videos?{params}", timeout=15) as r:
        data = json.loads(r.read().decode("utf-8"))
    if not data.get("items"):
        raise ValueError("Vídeo não encontrado")
    item = data["items"][0]
    stats = item.get("statistics", {})
    return {
        "views": int(stats.get("viewCount", 0)),
        "likes": int(stats.get("likeCount", 0)),
        "data": pd.Timestamp(item["snippet"]["publishedAt"][:10]),
    }


# --------------------------------------------------------------------------- #
# Análise por período
# --------------------------------------------------------------------------- #
def videos_vazio():
    df = pd.DataFrame({
        "Link": pd.Series(dtype="str"),
        "Embaixador": pd.Series(dtype="str"),
        "Rede": pd.Series(dtype="str"),
        "Data": pd.Series(dtype="datetime64[ns]"),
        "Visualizações": pd.Series(dtype="int64"),
        "Curtidas": pd.Series(dtype="int64"),
        "Salvamentos": pd.Series(dtype="int64"),
    })
    return df


def normalizar_videos(df):
    d = df.copy()
    for c in COLS_VIDEOS:
        if c not in d.columns:
            d[c] = None
    d = d[COLS_VIDEOS]
    d["Data"] = pd.to_datetime(d["Data"], errors="coerce")
    for c in COLS_NUM:
        d[c] = pd.to_numeric(d[c], errors="coerce").fillna(0).astype("int64")
    return d


def filtrar_periodo(df, ini, fim, redes=None):
    d = normalizar_videos(df)
    ini_ts, fim_ts = pd.Timestamp(ini), pd.Timestamp(fim) + pd.Timedelta(days=1)
    d = d[(d["Data"] >= ini_ts) & (d["Data"] < fim_ts)]
    if redes:
        d = d[d["Rede"].isin(redes)]
    return d.sort_values("Data").reset_index(drop=True)


def totais(df):
    return {
        "Vídeos": int(len(df)),
        "Visualizações": int(df["Visualizações"].sum()) if len(df) else 0,
        "Curtidas": int(df["Curtidas"].sum()) if len(df) else 0,
        "Salvamentos": int(df["Salvamentos"].sum()) if len(df) else 0,
    }


def resumo_por_rede(df):
    if df.empty:
        return pd.DataFrame(columns=["Rede", "Vídeos"] + COLS_NUM)
    g = df.groupby("Rede", dropna=False).agg(
        Vídeos=("Link", "size"),
        Visualizações=("Visualizações", "sum"),
        Curtidas=("Curtidas", "sum"),
        Salvamentos=("Salvamentos", "sum"),
    ).reset_index()
    return g.sort_values("Visualizações", ascending=False).reset_index(drop=True)


def fmt(n):
    return f"{int(n):,}".replace(",", ".")


# --------------------------------------------------------------------------- #
# Relatórios
# --------------------------------------------------------------------------- #
def gerar_excel(links_df, videos_df, ini, fim):
    from openpyxl.styles import Alignment, Font, PatternFill

    resumo = pd.DataFrame(
        [("Período", f"{ini:%d/%m/%Y} a {fim:%d/%m/%Y}")] + [(k, v) for k, v in totais(videos_df).items()],
        columns=["Indicador", "Valor"],
    )
    por_rede = resumo_por_rede(videos_df)
    videos_out = videos_df.copy()
    videos_out["Data"] = videos_out["Data"].dt.strftime("%d/%m/%Y")
    links_out = links_df if len(links_df) else pd.DataFrame(
        columns=["Original", "Rede", "Link corrigido", "Status", "Observação"])

    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as w:
        resumo.to_excel(w, sheet_name="Resumo", index=False)
        por_rede.to_excel(w, sheet_name="Por rede", index=False)
        videos_out.to_excel(w, sheet_name="Vídeos", index=False)
        links_out.to_excel(w, sheet_name="Links corrigidos", index=False)
        header_fill = PatternFill("solid", fgColor="5B3FD9")
        for ws in w.book.worksheets:
            for cell in ws[1]:
                cell.fill = header_fill
                cell.font = Font(bold=True, color="FFFFFF")
                cell.alignment = Alignment(horizontal="center", vertical="center")
            for col in ws.columns:
                largura = max((len(str(c.value)) if c.value is not None else 0) for c in col)
                ws.column_dimensions[col[0].column_letter].width = min(max(14, largura + 3), 60)
            ws.freeze_panes = "A2"
            for row in ws.iter_rows(min_row=2):
                for cell in row:
                    if isinstance(cell.value, int):
                        cell.number_format = "#,##0"
    return buf.getvalue()


def _t(s):
    return str(s).encode("latin-1", "replace").decode("latin-1")


def gerar_pdf(videos_df, ini, fim):
    from fpdf import FPDF

    class PDF(FPDF):
        def header(self):
            self.set_fill_color(18, 16, 62)
            self.rect(0, 0, 210, 26, "F")
            self.set_fill_color(255, 95, 184)
            self.rect(0, 26, 210, 1.2, "F")
            self.set_text_color(255, 255, 255)
            self.set_font("Helvetica", "B", 15)
            self.set_xy(10, 8)
            self.cell(0, 8, "BS LATAM | Programa de Embaixadores")
            self.set_y(33)
            self.set_text_color(30, 30, 70)

        def footer(self):
            self.set_y(-12)
            self.set_font("Helvetica", "I", 8)
            self.set_text_color(120, 120, 150)
            self.cell(0, 8, _t(f"Gerado em {datetime.now():%d/%m/%Y %H:%M}  |  Página {self.page_no()}"), align="C")

    pdf = PDF(format="A4")
    pdf.set_auto_page_break(auto=True, margin=16)
    pdf.add_page()

    pdf.set_font("Helvetica", "B", 17)
    pdf.cell(0, 10, _t("Relatório de desempenho de vídeos"), new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Helvetica", "", 11)
    pdf.set_text_color(90, 70, 190)
    pdf.cell(0, 7, _t(f"Período: {ini:%d/%m/%Y} a {fim:%d/%m/%Y}"), new_x="LMARGIN", new_y="NEXT")
    pdf.ln(4)

    # Cartões de totais
    t = totais(videos_df)
    x0, y0, w, h = 10, pdf.get_y(), 44, 20
    for i, (k, v) in enumerate(t.items()):
        x = x0 + i * (w + 4.67)
        pdf.set_fill_color(240, 236, 255)
        pdf.set_draw_color(139, 92, 246)
        pdf.rect(x, y0, w, h, "DF")
        pdf.set_xy(x, y0 + 3)
        pdf.set_font("Helvetica", "", 8)
        pdf.set_text_color(100, 90, 160)
        pdf.cell(w, 4, _t(k), align="C")
        pdf.set_xy(x, y0 + 9)
        pdf.set_font("Helvetica", "B", 13)
        pdf.set_text_color(30, 30, 90)
        pdf.cell(w, 8, fmt(v), align="C")
    pdf.set_y(y0 + h + 8)

    def tabela(titulo, colunas, linhas):
        pdf.set_font("Helvetica", "B", 12)
        pdf.set_text_color(30, 30, 70)
        pdf.cell(0, 8, _t(titulo), new_x="LMARGIN", new_y="NEXT")
        pdf.set_fill_color(91, 63, 217)
        pdf.set_text_color(255, 255, 255)
        pdf.set_font("Helvetica", "B", 9)
        for nome, larg in colunas:
            pdf.cell(larg, 7, _t(nome), border=0, fill=True, align="C")
        pdf.ln()
        pdf.set_font("Helvetica", "", 8.5)
        pdf.set_text_color(40, 40, 80)
        for i, linha in enumerate(linhas):
            pdf.set_fill_color(247, 244, 255) if i % 2 == 0 else pdf.set_fill_color(255, 255, 255)
            for (nome, larg), val in zip(colunas, linha):
                numerico = bool(re.match(r"^[\d.]+$", str(val)))
                pdf.cell(larg, 6.5, _t(val), fill=True, align="R" if numerico else "L")
            pdf.ln()
        pdf.ln(5)

    por_rede = resumo_por_rede(videos_df)
    tabela(
        "Por rede social",
        [("Rede", 50), ("Vídeos", 30), ("Visualizações", 40), ("Curtidas", 35), ("Salvamentos", 35)],
        [(r["Rede"], fmt(r["Vídeos"]), fmt(r["Visualizações"]), fmt(r["Curtidas"]), fmt(r["Salvamentos"]))
         for _, r in por_rede.iterrows()],
    )

    linhas = []
    for _, r in videos_df.iterrows():
        link = str(r["Link"])
        link = link if len(link) <= 30 else link[:27] + "..."
        data = r["Data"].strftime("%d/%m/%Y") if pd.notna(r["Data"]) else "-"
        linhas.append((str(r["Embaixador"] or "-")[:20], str(r["Rede"] or "-"), data, link,
                       fmt(r["Visualizações"]), fmt(r["Curtidas"]), fmt(r["Salvamentos"])))
    tabela(
        "Vídeos do período",
        [("Embaixador", 32), ("Rede", 20), ("Data", 22), ("Link", 44), ("Views", 24), ("Curtidas", 24),
         ("Salv.", 24)],
        linhas,
    )
    return bytes(pdf.output())
